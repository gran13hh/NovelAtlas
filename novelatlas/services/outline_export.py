"""Stateless HTML and DOCX renderers for validated novel outlines."""

import re
from datetime import UTC, datetime
from html import escape
from io import BytesIO

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from novelatlas.schemas.analysis import NovelOutline, OutlineSource
from novelatlas.schemas.exports import ExportSection

_GREEN = RGBColor(49, 83, 63)
_DARK_GREEN = RGBColor(30, 50, 39)
_MUTED = RGBColor(101, 105, 99)
_BODY = RGBColor(32, 34, 31)
_FONT = "Arial Unicode MS"

SECTION_TITLES: dict[ExportSection, str] = {
    "overall_summary": "总体概述",
    "chapter_outline": "章节范围细纲",
    "storylines": "主要故事线",
    "characters": "主要人物与关系",
    "worldbuilding": "世界观",
    "foreshadowing": "伏笔",
    "unresolved_items": "未解决事件与不确定项",
}


def safe_export_filename(title: str, extension: str) -> str:
    """Build a human-readable filename without path or header characters."""

    stem = re.sub(r"\.txt$", "", title, flags=re.IGNORECASE)
    stem = re.sub(r"[<>:\"/\\|?*\x00-\x1f]", "_", stem).strip(" ._")
    if not stem:
        stem = "NovelAtlas"
    return f"{stem[:120]}-全书细纲.{extension}"


def _set_run_font(
    run,
    *,
    size: float | None = None,
    color: RGBColor | None = None,
    bold: bool | None = None,
    italic: bool | None = None,
) -> None:
    run.font.name = _FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), _FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), _FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), _FONT)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def _configure_styles(document: Document) -> None:
    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = _FONT
    normal.font.size = Pt(11)
    normal.font.color.rgb = _BODY
    normal._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), _FONT)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for name, size, before, after, color in (
        ("Heading 1", 16, 18, 10, _DARK_GREEN),
        ("Heading 2", 13, 14, 7, _GREEN),
        ("Heading 3", 12, 10, 5, _GREEN),
    ):
        style = styles[name]
        style.font.name = _FONT
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), _FONT)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    bullet = styles["List Bullet"]
    bullet.font.name = _FONT
    bullet.font.size = Pt(11)
    bullet.font.color.rgb = _BODY
    bullet._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), _FONT)
    bullet.paragraph_format.left_indent = Inches(0.375)
    bullet.paragraph_format.first_line_indent = Inches(-0.188)
    bullet.paragraph_format.space_after = Pt(4)
    bullet.paragraph_format.line_spacing = 1.25

    if "Source Note" not in styles:
        source_style = styles.add_style("Source Note", WD_STYLE_TYPE.PARAGRAPH)
    else:
        source_style = styles["Source Note"]
    source_style.font.name = _FONT
    source_style.font.size = Pt(8.5)
    source_style.font.color.rgb = _MUTED
    source_style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), _FONT)
    source_style.paragraph_format.space_before = Pt(1)
    source_style.paragraph_format.space_after = Pt(6)


def _add_field(paragraph, instruction: str, placeholder: str = "") -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction_element = OxmlElement("w:instrText")
    instruction_element.set(qn("xml:space"), "preserve")
    instruction_element.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    elements = [begin, instruction_element, separate]
    for index, line in enumerate(placeholder.split("\n")):
        if index:
            elements.append(OxmlElement("w:br"))
        text = OxmlElement("w:t")
        text.text = line
        elements.append(text)
    elements.append(end)
    run._r.extend(elements)


def _configure_document(document: Document, title: str) -> None:
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    section.different_first_page_header_footer = True

    header = section.header
    header_paragraph = header.paragraphs[0]
    header_paragraph.text = "NovelAtlas · 全书细纲"
    header_paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _set_run_font(header_paragraph.runs[0], size=8.5, color=_MUTED)

    footer = section.footer
    footer_paragraph = footer.paragraphs[0]
    footer_paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    prefix = footer_paragraph.add_run("第 ")
    _set_run_font(prefix, size=8.5, color=_MUTED)
    _add_field(footer_paragraph, " PAGE ", "1")
    suffix = footer_paragraph.add_run(" 页")
    _set_run_font(suffix, size=8.5, color=_MUTED)

    settings = document.settings.element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")

    document.core_properties.title = f"{title} - 全书细纲"
    document.core_properties.subject = "NovelAtlas 生成的长篇小说全书细纲"
    document.core_properties.author = "NovelAtlas"
    document.core_properties.keywords = "小说, 细纲, 故事线, 人物关系, 世界观"


def _add_cover(document: Document, title: str, generated_at: datetime) -> None:
    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(96)

    kicker = document.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    kicker.paragraph_format.space_after = Pt(18)
    _set_run_font(
        kicker.add_run("NOVELATLAS · 全书细纲"),
        size=10,
        color=_GREEN,
        bold=True,
    )

    title_paragraph = document.add_paragraph()
    title_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_paragraph.paragraph_format.space_after = Pt(12)
    _set_run_font(
        title_paragraph.add_run(title),
        size=28,
        color=_DARK_GREEN,
        bold=True,
    )

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(100)
    _set_run_font(
        subtitle.add_run("故事线 · 人物关系 · 世界观 · 伏笔与未解事项"),
        size=12,
        color=_MUTED,
    )

    date_paragraph = document.add_paragraph()
    date_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_run_font(
        date_paragraph.add_run(generated_at.strftime("%Y 年 %m 月 %d 日")),
        size=10,
        color=_MUTED,
    )
    document.add_page_break()


def _add_toc(document: Document, sections: list[ExportSection]) -> None:
    document.add_heading("目录", level=1)
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(10)
    placeholder = "\n".join(SECTION_TITLES[section] for section in sections)
    _add_field(paragraph, ' TOC \\o "1-2" \\h \\z \\u ', placeholder)
    document.add_page_break()


def _source_note(source: OutlineSource) -> str:
    chapter_count = len(set(source.chapter_ids))
    batch_count = len(set(source.batch_ids))
    return f"来源覆盖：{chapter_count} 个章节引用 · {batch_count} 个分析批次"


def _add_bullet(document: Document, text: str) -> None:
    paragraph = document.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.left_indent = Inches(0.375)
    paragraph.paragraph_format.first_line_indent = Inches(-0.188)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing = 1.25
    paragraph.add_run(text)


def _add_labeled_paragraph(document: Document, label: str, text: str) -> None:
    paragraph = document.add_paragraph()
    label_run = paragraph.add_run(f"{label}：")
    _set_run_font(label_run, bold=True, color=_GREEN)
    paragraph.add_run(text)


def _add_outline_sections(
    document: Document,
    outline: NovelOutline,
    sections: list[ExportSection],
) -> None:
    if "overall_summary" in sections:
        document.add_heading(SECTION_TITLES["overall_summary"], level=1)
        document.add_paragraph(outline.overall_summary)

    if "chapter_outline" in sections:
        document.add_heading(SECTION_TITLES["chapter_outline"], level=1)
        if not outline.chapter_outline:
            document.add_paragraph("暂无明确的章节范围细纲。")
        for item in outline.chapter_outline:
            document.add_heading(item.chapter_range, level=2)
            document.add_paragraph(item.summary)
            for event in item.key_events:
                _add_bullet(document, event)
            document.add_paragraph(_source_note(item.sources), style="Source Note")

    if "storylines" in sections:
        document.add_heading(SECTION_TITLES["storylines"], level=1)
        if not outline.storylines:
            document.add_paragraph("暂无明确的主要故事线。")
        for item in outline.storylines:
            document.add_heading(item.name, level=2)
            document.add_paragraph(item.summary)
            for development in item.developments:
                _add_bullet(document, development)
            document.add_paragraph(_source_note(item.sources), style="Source Note")

    if "characters" in sections:
        document.add_heading(SECTION_TITLES["characters"], level=1)
        if not outline.characters:
            document.add_paragraph("暂无明确的主要人物归纳。")
        for item in outline.characters:
            document.add_heading(item.name, level=2)
            document.add_paragraph(item.summary)
            for relationship in item.relationships:
                _add_labeled_paragraph(document, "关系", relationship)
            for change in item.changes:
                _add_labeled_paragraph(document, "变化", change)
            document.add_paragraph(_source_note(item.sources), style="Source Note")

    if "worldbuilding" in sections:
        document.add_heading(SECTION_TITLES["worldbuilding"], level=1)
        if not outline.worldbuilding:
            document.add_paragraph("暂无明确的世界观条目。")
        for item in outline.worldbuilding:
            document.add_heading(item.name, level=2)
            _add_labeled_paragraph(document, "类别", item.category)
            document.add_paragraph(item.description)
            document.add_paragraph(_source_note(item.sources), style="Source Note")

    claim_sections = (
        ("foreshadowing", outline.foreshadowing),
        (
            "unresolved_items",
            [
                *outline.unresolved_items,
                *outline.conflicts_and_uncertainties,
            ],
        ),
    )
    for section_name, claims in claim_sections:
        if section_name not in sections:
            continue
        typed_section = section_name  # Help type checkers retain the literal mapping.
        document.add_heading(SECTION_TITLES[typed_section], level=1)  # type: ignore[index]
        if not claims:
            document.add_paragraph("暂无明确内容。")
        for claim in claims:
            suffix = "" if claim.confidence == "certain" else f"（{claim.confidence}）"
            _add_bullet(document, f"{claim.description}{suffix}")
            document.add_paragraph(_source_note(claim.sources), style="Source Note")


def render_outline_docx(
    *,
    title: str,
    outline: NovelOutline,
    sections: list[ExportSection],
) -> bytes:
    """Render a polished Word document entirely in memory."""

    document = Document()
    _configure_styles(document)
    _configure_document(document, title)
    _add_cover(document, title, datetime.now(UTC))
    _add_toc(document, sections)
    _add_outline_sections(document, outline, sections)

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _html_source_note(source: OutlineSource) -> str:
    return (
        f"来源覆盖：{len(set(source.chapter_ids))} 个章节引用 · "
        f"{len(set(source.batch_ids))} 个分析批次"
    )


def render_outline_html(
    *,
    title: str,
    outline: NovelOutline,
    sections: list[ExportSection],
) -> bytes:
    """Render a self-contained, escaped, print-friendly HTML report."""

    safe_title = escape(title)
    section_ids: dict[ExportSection, str] = {
        section: f"section-{section.replace('_', '-')}" for section in sections
    }
    toc = "".join(
        f'<a href="#{section_ids[section]}">{escape(SECTION_TITLES[section])}</a>'
        for section in sections
    )
    body: list[str] = []

    def open_section(section: ExportSection) -> None:
        body.append(
            f'<section id="{section_ids[section]}"><h2>{escape(SECTION_TITLES[section])}</h2>'
        )

    if "overall_summary" in sections:
        open_section("overall_summary")
        body.append(f"<p>{escape(outline.overall_summary)}</p></section>")

    if "chapter_outline" in sections:
        open_section("chapter_outline")
        if not outline.chapter_outline:
            body.append('<p class="empty">暂无明确的章节范围细纲。</p>')
        for item in outline.chapter_outline:
            events = "".join(f"<li>{escape(event)}</li>" for event in item.key_events)
            body.append(
                f'<article><h3>{escape(item.chapter_range)}</h3><p>{escape(item.summary)}</p>'
                f'{f"<ul>{events}</ul>" if events else ""}'
                f'<p class="source">{escape(_html_source_note(item.sources))}</p></article>'
            )
        body.append("</section>")

    if "storylines" in sections:
        open_section("storylines")
        if not outline.storylines:
            body.append('<p class="empty">暂无明确的主要故事线。</p>')
        for item in outline.storylines:
            developments = "".join(
                f"<li>{escape(value)}</li>" for value in item.developments
            )
            body.append(
                f'<article><h3>{escape(item.name)}</h3><p>{escape(item.summary)}</p>'
                f'{f"<ul>{developments}</ul>" if developments else ""}'
                f'<p class="source">{escape(_html_source_note(item.sources))}</p></article>'
            )
        body.append("</section>")

    if "characters" in sections:
        open_section("characters")
        if not outline.characters:
            body.append('<p class="empty">暂无明确的主要人物归纳。</p>')
        for item in outline.characters:
            details = "".join(
                f"<li><strong>关系：</strong>{escape(value)}</li>"
                for value in item.relationships
            ) + "".join(
                f"<li><strong>变化：</strong>{escape(value)}</li>"
                for value in item.changes
            )
            body.append(
                f'<article><h3>{escape(item.name)}</h3><p>{escape(item.summary)}</p>'
                f'{f"<ul>{details}</ul>" if details else ""}'
                f'<p class="source">{escape(_html_source_note(item.sources))}</p></article>'
            )
        body.append("</section>")

    if "worldbuilding" in sections:
        open_section("worldbuilding")
        if not outline.worldbuilding:
            body.append('<p class="empty">暂无明确的世界观条目。</p>')
        for item in outline.worldbuilding:
            body.append(
                f'<article><h3>{escape(item.name)}</h3>'
                f'<p class="tag">{escape(item.category)}</p>'
                f'<p>{escape(item.description)}</p>'
                f'<p class="source">{escape(_html_source_note(item.sources))}</p></article>'
            )
        body.append("</section>")

    for section_name, claims in (
        ("foreshadowing", outline.foreshadowing),
        (
            "unresolved_items",
            [*outline.unresolved_items, *outline.conflicts_and_uncertainties],
        ),
    ):
        if section_name not in sections:
            continue
        typed_section: ExportSection = section_name  # type: ignore[assignment]
        open_section(typed_section)
        if not claims:
            body.append('<p class="empty">暂无明确内容。</p>')
        else:
            items = "".join(
                f'<li>{escape(claim.description)}'
                f'<span class="confidence">{escape(claim.confidence)}</span>'
                f'<small>{escape(_html_source_note(claim.sources))}</small></li>'
                for claim in claims
            )
            body.append(f'<ul class="claims">{items}</ul>')
        body.append("</section>")

    generated = datetime.now(UTC).strftime("%Y-%m-%d %H:%M UTC")
    html = f"""<!doctype html>
<html lang="zh-CN">
<head>
  <meta charset="utf-8">
  <meta name="viewport" content="width=device-width, initial-scale=1">
  <meta name="color-scheme" content="light">
  <title>{safe_title} - 全书细纲</title>
  <style>
    :root {{ color: #20221f; background: #f3f1ea; font-family: "PingFang SC", "Microsoft YaHei", sans-serif; }}
    * {{ box-sizing: border-box; }}
    body {{ margin: 0; line-height: 1.8; }}
    main {{ width: min(920px, calc(100% - 32px)); margin: 32px auto; background: #fbfaf6; padding: clamp(28px, 6vw, 72px); box-shadow: 0 18px 50px rgba(30,50,39,.10); }}
    header {{ padding: 48px 0 56px; border-bottom: 1px solid #d7d6cf; }}
    .kicker {{ color: #55705e; font-size: 12px; font-weight: 700; letter-spacing: .18em; text-transform: uppercase; }}
    h1 {{ margin: 14px 0 12px; color: #1e3227; font-family: Georgia, "Songti SC", serif; font-size: clamp(36px, 7vw, 60px); line-height: 1.15; }}
    .meta, .source, small {{ color: #74766f; font-size: 12px; }}
    nav {{ display: flex; flex-wrap: wrap; gap: 8px; margin: 28px 0 48px; }}
    nav a {{ color: #31533f; background: #e4ebe0; border-radius: 999px; padding: 6px 13px; text-decoration: none; font-size: 13px; font-weight: 600; }}
    section {{ margin-top: 50px; scroll-margin-top: 20px; }}
    h2 {{ color: #1e3227; font-family: Georgia, "Songti SC", serif; font-size: 28px; border-bottom: 1px solid #d7d6cf; padding-bottom: 9px; }}
    h3 {{ color: #31533f; margin: 28px 0 7px; font-size: 18px; }}
    p {{ white-space: pre-wrap; }}
    article {{ break-inside: avoid; }}
    ul {{ padding-left: 1.3em; }}
    li {{ margin: 7px 0; }}
    .claims {{ list-style: none; padding: 0; }}
    .claims li {{ padding: 14px 16px; border: 1px solid #deddd6; border-radius: 10px; }}
    .claims small {{ display: block; margin-top: 4px; }}
    .confidence, .tag {{ display: inline-block; margin-left: 8px; color: #31533f; background: #e4ebe0; border-radius: 999px; padding: 1px 8px; font-size: 11px; }}
    .empty {{ color: #777970; font-style: italic; }}
    footer {{ margin-top: 64px; padding-top: 16px; border-top: 1px solid #d7d6cf; color: #777970; font-size: 12px; }}
    @media print {{ :root {{ background: white; }} main {{ width: auto; margin: 0; padding: 0; box-shadow: none; }} nav {{ display: none; }} section {{ break-before: auto; }} }}
  </style>
</head>
<body>
  <main>
    <header><p class="kicker">NovelAtlas · 全书细纲</p><h1>{safe_title}</h1><p class="meta">生成时间：{generated}</p></header>
    <nav aria-label="细纲目录">{toc}</nav>
    {''.join(body)}
    <footer>由 NovelAtlas 根据已确认的结构化细纲生成。</footer>
  </main>
</body>
</html>"""
    return html.encode("utf-8")
